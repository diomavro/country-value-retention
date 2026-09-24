"""Parse CYSTAT PxWeb exports (layouts documented in data/raw/cystat/STRUCTURE.md).

Output: long-format Parquet with one row per published cell and the provenance
columns required by ``cvr.provenance``.  Confidential/not-applicable symbols
are kept as ``flag`` with ``value`` = NaN; they are never read as zero.
"""

from __future__ import annotations

import re
from pathlib import Path

import openpyxl
import pandas as pd

from ..provenance import require, retrieval_date

RAW = Path("data/raw/cystat")
OUT = Path("data/interim/cystat")
CODE = re.compile(r"^\[([^\]]+)\]")
SYMBOLS = {"c": "confidential", "C": "confidential", "N.A.": "not_applicable", "...": "not_available"}
FIRST_YEAR, LAST_YEAR = 2010, 2022
PXWEB = "https://cystatdb.cystat.gov.cy/pxweb/en/8.CYSTAT-DB/"

SUIOT_TABLES = {
    "0640010E": "Supply table at basic prices incl. transformation to purchasers' prices",
    "0640015E": "Use table at purchasers' prices",
    "0640020E": "Use table at basic prices",
    "0640025E": "Use table for domestic output at basic prices",
    "0640030E": "Use table for imports at basic prices",
    "0640035E": "Trade and transport margins",
    "0640040E": "Taxes less subsidies on products",
    "0640045E": "Symmetric IOT product x product, total",
    "0640050E": "Symmetric IOT product x product, domestic output",
    "0640055E": "Symmetric IOT product x product, imports",
}


def _code(label: object) -> str | None:
    if label is None:
        return None
    m = CODE.match(str(label).strip())
    return m.group(1).replace("P3_S14", "P3-S14") if m else None


def _manifest() -> pd.DataFrame:
    return pd.read_csv(RAW / "MANIFEST.csv")


def parse_suiot(table: str, measure: str = "CP") -> pd.DataFrame:
    path = RAW / f"{table}_{measure}.xlsx"
    ws = openpyxl.load_workbook(path, read_only=True, data_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    header = rows[2]
    cols = {j: _code(h) for j, h in enumerate(header) if j >= 3 and _code(h)}

    # Block length B = distance between consecutive year markers in column A.
    starts = [
        i
        for i, r in enumerate(rows)
        if i >= 3 and r[0] is not None and str(r[0]).strip().isdigit()
    ]
    if len(starts) != LAST_YEAR - FIRST_YEAR + 1:
        raise ValueError(
            f"{table}: expected {LAST_YEAR - FIRST_YEAR + 1} year blocks, found {len(starts)}"
        )
    B = starts[1] - starts[0]
    if any(b - a != B for a, b in zip(starts, starts[1:])):
        raise ValueError(f"{table}: irregular year blocks")

    recs = []
    for s in starts:
        year = int(str(rows[s][0]).strip())
        for i in range(s, s + B):
            rcode = _code(rows[i][2])
            if rcode is None:
                raise ValueError(f"{table} {year}: unlabeled row {i + 1}")
            for j, ccode in cols.items():
                v = rows[i][j] if j < len(rows[i]) else None
                if isinstance(v, (int, float)):
                    recs.append((year, rcode, ccode, float(v), ""))
                else:
                    sv = "" if v is None else str(v).strip()
                    if sv not in SYMBOLS:
                        raise ValueError(
                            f"{table} {year} {rcode}/{ccode}: unexpected cell {v!r}"
                        )
                    recs.append((year, rcode, ccode, float("nan"), SYMBOLS[sv]))
    df = pd.DataFrame(
        recs, columns=["reference_year", "row_code", "col_code", "value", "flag"]
    )
    man = _manifest()
    hit = man[man["file"].astype(str).str.contains(f"{table}_{measure}")]
    df.insert(0, "table_id", table)
    df["measure"] = measure
    df["unit"] = "EUR million"
    df["source"] = f"CYSTAT SUIOT {table}: {SUIOT_TABLES[table]}"
    df["source_url"] = hit["url"].iloc[0] if len(hit) else PXWEB
    df["retrieval_date"] = retrieval_date(f"{table}_{measure}.xlsx")
    df["methodology"] = "CYSTAT PxWeb xlsx export; cell values as published"
    df["confidence_level"] = "high"
    df["status"] = "observed"
    return require(df, table)


def parse_sector_accounts() -> pd.DataFrame:
    """0630030E non-financial sector accounts, 2010–2024, incl. S2 rest of world."""
    ws = openpyxl.load_workbook(
        RAW / "0630030E.xlsx", read_only=True, data_only=True
    ).active
    rows = list(ws.iter_rows(values_only=True))
    header = rows[2]
    recs, year = [], None
    for r in rows[3:]:
        if r[0] is not None and str(r[0]).strip().isdigit():
            year = int(str(r[0]).strip())
        if r[1] is None or year is None:
            continue
        sector = str(r[1]).split()[0].strip()
        if not re.fullmatch(r"S\d+", sector):
            break
        for j in range(2, len(header)):
            h = header[j]
            if h is None:
                continue
            item = str(h).split()[0]
            v = r[j]
            if isinstance(v, (int, float)):
                recs.append((year, sector, item, str(h), float(v), ""))
            else:
                sv = "" if v is None else str(v).strip()
                recs.append(
                    (
                        year,
                        sector,
                        item,
                        str(h),
                        float("nan"),
                        SYMBOLS.get(sv, sv or "blank"),
                    )
                )
    df = pd.DataFrame(
        recs,
        columns=["reference_year", "sector", "item", "item_label", "value", "flag"],
    )
    df["unit"] = "EUR million"
    df["source"] = "CYSTAT 0630030E non-financial institutional sector accounts"
    df["source_url"] = PXWEB
    df["retrieval_date"] = retrieval_date("0630030E.xlsx")
    df["methodology"] = "CYSTAT PxWeb xlsx export; cell values as published"
    df["confidence_level"] = "high"
    df["status"] = "observed"
    return require(df, "0630030E")


def parse_gdp_gni() -> pd.DataFrame:
    """0610010E: GDP and GNI (B5G) at current prices, 1995–2025."""
    ws = openpyxl.load_workbook(
        RAW / "0610010E.xlsx", read_only=True, data_only=True
    ).active
    rows = list(ws.iter_rows(values_only=True))
    years = [
        int(str(y)) for y in rows[2][1:] if y is not None and str(y).strip().isdigit()
    ]
    recs = []
    for r in rows[3:16]:
        label = str(r[0])
        for y, v in zip(years, r[1:]):
            recs.append(
                (y, label, float(v) if isinstance(v, (int, float)) else float("nan"))
            )
    df = pd.DataFrame(recs, columns=["reference_year", "item_label", "value"])
    df["flag"] = df["value"].isna().map({True: "not_available", False: ""})
    df["source"] = "CYSTAT 0610010E GDP and GNI"
    df["source_url"] = PXWEB
    df["retrieval_date"] = retrieval_date("0610010E.xlsx")
    df["methodology"] = "CYSTAT PxWeb xlsx export; cell values as published"
    df["confidence_level"] = "high"
    df["status"] = "observed"
    return require(df, "0610010E")


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    frames = [parse_suiot(t) for t in SUIOT_TABLES]
    suiot = pd.concat(frames, ignore_index=True)
    suiot.to_parquet(OUT / "suiot_cp.parquet", index=False)
    parse_sector_accounts().to_parquet(OUT / "sector_accounts.parquet", index=False)
    parse_gdp_gni().to_parquet(OUT / "gdp_gni.parquet", index=False)
    print(
        f"SUIOT cells: {len(suiot):,} across {suiot.table_id.nunique()} tables, years {suiot.reference_year.min()}–{suiot.reference_year.max()}"
    )


if __name__ == "__main__":
    main()
